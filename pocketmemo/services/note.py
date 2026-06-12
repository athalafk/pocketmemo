"""Note service — titled notes with optional file attachments and export.

A note has a title + text body, is searchable via embedding, and may have photo
or document attachments (files whose note_id points to it). Notes can be updated,
exported to .docx/.txt, and have handwritten image content read into them (OCR).
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from sqlalchemy import select
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.ext import ContextTypes

from pocketmemo.database import SessionLocal
from pocketmemo.i18n import t
from pocketmemo.llm import llm
from pocketmemo.models import Folder, Note, StoredFile, User

NOTE_FOLDER_MATCH_SYSTEM_PROMPT = (
    "You sort a note into the best-fitting folder. Answer with ONLY a single "
    "number: the folder's number, or 0 if none fits well. No other text."
)

logger = logging.getLogger(__name__)

RECALL_CANDIDATES = 5
RECALL_PREFILTER_MAX_DISTANCE = 0.9
EXPORT_ROOT = Path("/app/storage/exports")

NOTE_MATCH_SYSTEM_PROMPT = (
    "You are a note matcher. Pick the ONE note from the list that best matches what "
    "the user is looking for, or say none. Answer with ONLY a single number: the "
    "matching note's number, or 0 if none is relevant. No other text."
)

UPDATE_NOTE_SYSTEM_PROMPT = (
    "You help update a note. Given the list of notes and the user's request, output "
    "JSON:\n"
    "- note_id: the id the user means (number), or 0 if none matches.\n"
    "- changes: an object with ONLY what to change:\n"
    "  - title: new title (string)\n"
    "  - content: text (string)\n"
    "  - append: true if the text is ADDED to the existing body, false to REPLACE it.\n"
    "Include only relevant fields. Output ONLY JSON."
)


def _parse_choice(raw: str, n: int) -> int | None:
    match = re.search(r"\d+", raw or "")
    if not match:
        return None
    idx = int(match.group())
    return idx if 1 <= idx <= n else None


def _safe_filename(name: str) -> str:
    name = re.sub(r"[^\w\- ]", "", name or "note", flags=re.UNICODE).strip()
    return name[:80] or "note"


async def save_note(user: User, title: str, content: str) -> str:
    """Save a new note (title + body)."""
    lang = user.language
    title = (title or "").strip()
    content = (content or "").strip()
    if not title and not content:
        return t("note_empty", lang)
    if not title:
        title = content[:60] + ("…" if len(content) > 60 else "")
    if not content:
        content = title

    embedding = await llm.embed(f"{title}\n{content}")
    async with SessionLocal() as session:
        note = Note(user_id=user.id, title=title, content=content, embedding=embedding)
        session.add(note)
        await session.commit()
        await session.refresh(note)

    logger.info("Saved note %s for user %s", note.id, user.id)
    reply = t("note_saved", lang, id=note.id, title=title)
    folder_name = await _auto_categorize_note(user, note.id, title, content)
    if folder_name:
        reply += "\n" + t("note_filed_into", lang, folder=folder_name)
    return reply


async def _auto_categorize_note(
    user: User, note_id: int, title: str, content: str
) -> str | None:
    """Let the LLM file a new note into the best matching note-folder, if any."""
    async with SessionLocal() as session:
        folders = list(
            (
                await session.execute(
                    select(Folder)
                    .where(Folder.user_id == user.id, Folder.kind == "note")
                    .order_by(Folder.name)
                )
            ).scalars().all()
        )
    if not folders:
        return None

    listing = "\n".join(f"{i}. {f.name}" for i, f in enumerate(folders, start=1))
    prompt = (
        f"Note title: {title}\nNote content: {content[:500]}\n\n"
        f"Folders:\n{listing}\n\n"
        "Which folder number fits best? Answer 0 if none fits."
    )
    try:
        raw = await llm.chat(prompt, system_prompt=NOTE_FOLDER_MATCH_SYSTEM_PROMPT)
    except Exception:
        logger.exception("Auto-categorize note failed")
        return None
    choice = _parse_choice(raw, len(folders))
    if choice is None:
        return None

    folder = folders[choice - 1]
    async with SessionLocal() as session:
        note = (
            await session.execute(
                select(Note).where(Note.id == note_id, Note.user_id == user.id)
            )
        ).scalar_one_or_none()
        if note is None:
            return None
        note.folder_id = folder.id
        await session.commit()
    logger.info("Auto-filed note %s into folder %s", note_id, folder.id)
    return folder.name


async def recall_note(
    update: Update, context: ContextTypes.DEFAULT_TYPE, user: User, query: str
) -> str | None:
    """Find the best note, send its text + attachments. Returns text on failure."""
    lang = user.language
    query = (query or "").strip()
    if not query:
        return t("note_recall_what", lang)

    query_embedding = await llm.embed_query(query)
    async with SessionLocal() as session:
        dist = Note.embedding.cosine_distance(query_embedding).label("dist")
        stmt = (
            select(Note, dist)
            .where(Note.user_id == user.id, Note.embedding.is_not(None))
            .order_by(dist)
            .limit(RECALL_CANDIDATES)
        )
        rows = (await session.execute(stmt)).all()

    if not rows:
        return t("note_recall_none", lang)

    candidates = [n for n, d in rows if d is None or d <= RECALL_PREFILTER_MAX_DISTANCE]
    if not candidates:
        return t("note_recall_no_match", lang, query=query)

    listing = "\n".join(f"{i}. {n.title}" for i, n in enumerate(candidates, start=1))
    prompt = (
        f"The user is looking for a note: \"{query}\".\n\n"
        f"Notes:\n{listing}\n\n"
        "Which number matches best? Answer 0 if none is relevant."
    )
    raw = await llm.chat(prompt, system_prompt=NOTE_MATCH_SYSTEM_PROMPT)
    choice = _parse_choice(raw, len(candidates))
    if choice is None:
        return t("note_recall_no_match", lang, query=query)

    note = candidates[choice - 1]
    files = await _note_attachments(note.id)
    await _send_note(context, update.effective_chat.id, note, files, lang)
    logger.info("Recalled note %s (%d attachments) for user %s", note.id, len(files), user.id)
    return None


async def _note_attachments(note_id: int) -> list[StoredFile]:
    async with SessionLocal() as session:
        return list(
            (
                await session.execute(
                    select(StoredFile).where(StoredFile.note_id == note_id).order_by(StoredFile.id)
                )
            ).scalars().all()
        )


async def _send_note(context, chat_id: int, note: Note, files: list[StoredFile], lang: str) -> None:
    keyboard = InlineKeyboardMarkup(
        [
            [
                InlineKeyboardButton("📄 .docx", callback_data=f"ndocx:{note.id}"),
                InlineKeyboardButton("📄 .txt", callback_data=f"ntxt:{note.id}"),
            ],
            [
                InlineKeyboardButton(t("note_btn_edit", lang), callback_data=f"nedit:{note.id}"),
                InlineKeyboardButton(t("note_btn_files", lang), callback_data=f"nfiles:{note.id}"),
            ],
        ]
    )
    await context.bot.send_message(
        chat_id=chat_id,
        text=f"📝 {note.title}\n\n{note.content or t('note_no_content', lang)}",
        reply_markup=keyboard,
    )
    for f in files:
        try:
            if f.file_type == "photo":
                await context.bot.send_photo(chat_id, photo=f.telegram_file_id, caption=f.display_name)
            else:
                await context.bot.send_document(chat_id, document=f.telegram_file_id, caption=f.display_name)
        except Exception:
            logger.exception("Failed to send attachment of note %s (file %s)", note.id, f.id)


async def view_note(context, user: User, chat_id: int, note_id: int) -> str | None:
    """Send a note (text + attachments) by id. Returns text on failure."""
    lang = user.language
    async with SessionLocal() as session:
        note = (
            await session.execute(
                select(Note).where(Note.id == note_id, Note.user_id == user.id)
            )
        ).scalar_one_or_none()
    if note is None:
        return t("note_not_found", lang, id=note_id)
    files = await _note_attachments(note_id)
    await _send_note(context, chat_id, note, files, lang)
    return None


async def get_active_notes(user: User) -> list[Note]:
    async with SessionLocal() as session:
        stmt = select(Note).where(Note.user_id == user.id).order_by(Note.created_at.desc())
        return list((await session.execute(stmt)).scalars().all())


async def get_recent_notes(user: User, limit: int = 3) -> list[Note]:
    async with SessionLocal() as session:
        stmt = (
            select(Note)
            .where(Note.user_id == user.id)
            .order_by(Note.created_at.desc())
            .limit(limit)
        )
        return list((await session.execute(stmt)).scalars().all())


def format_note_line(note: Note) -> str:
    return f"#{note.id} 📝 {note.title}"


async def delete_note(user: User, note_id: int) -> str:
    async with SessionLocal() as session:
        note = (
            await session.execute(
                select(Note).where(Note.id == note_id, Note.user_id == user.id)
            )
        ).scalar_one_or_none()
        if note is None:
            return t("note_not_found", user.language, id=note_id)
        await session.delete(note)
        await session.commit()
    logger.info("Deleted note %s for user %s", note_id, user.id)
    return t("note_deleted", user.language, id=note_id)


async def attach_file_to_note(
    user: User, file_id: int, note_id: int
) -> tuple[str, bool]:
    """Link a file to a note. Returns (reply_text, can_ocr)."""
    async with SessionLocal() as session:
        f = (
            await session.execute(
                select(StoredFile).where(
                    StoredFile.id == file_id, StoredFile.user_id == user.id
                )
            )
        ).scalar_one_or_none()
        note = (
            await session.execute(
                select(Note).where(Note.id == note_id, Note.user_id == user.id)
            )
        ).scalar_one_or_none()
        if f is None or note is None:
            return t("note_attach_failed", user.language), False
        f.note_id = note_id
        can_ocr = f.file_type == "photo" or f.mime_type == "application/pdf"
        await session.commit()
        title = note.title

    logger.info("Attached file %s to note %s (user %s)", file_id, note_id, user.id)
    return t("note_attached", user.language, title=title), can_ocr


async def transcribe_into_note(user: User, file_id: int, note_id: int) -> str:
    """Read an attached image/PDF via vision and append it under the note body."""
    lang = user.language
    async with SessionLocal() as session:
        f = (
            await session.execute(
                select(StoredFile).where(
                    StoredFile.id == file_id, StoredFile.user_id == user.id
                )
            )
        ).scalar_one_or_none()
        if f is None:
            return t("note_transcribe_filenotfound", lang)
        file_path = f.file_path
        mime = f.mime_type or "image/jpeg"

    try:
        data = Path(file_path).read_bytes()
    except OSError:
        return t("note_transcribe_filemissing", lang)

    text = await llm.transcribe_media(data, mime)
    if not text:
        return t("note_transcribe_failed", lang)

    async with SessionLocal() as session:
        note = (
            await session.execute(
                select(Note).where(Note.id == note_id, Note.user_id == user.id)
            )
        ).scalar_one_or_none()
        if note is None:
            return t("note_not_found", lang, id=note_id)
        existing = note.content or ""
        header = t("note_section_from_image", lang)
        note.content = f"{existing}\n\n{header}\n{text}".strip()
        note.embedding = await llm.embed(f"{note.title}\n{note.content}")
        await session.commit()

    logger.info("Transcribed file %s into note %s", file_id, note_id)
    return t("note_transcribe_done", lang)


# ---- Update ---------------------------------------------------------------

async def _parse_note_update(text: str, notes: list[Note]) -> dict:
    listing = "\n".join(f"#{n.id} {n.title}" for n in notes)
    prompt = (
        f"Notes:\n{listing}\n\n"
        f"User request: \"{text}\"\n"
        "Output JSON {note_id, changes} as instructed."
    )
    return await llm.complete_json(prompt, system_prompt=UPDATE_NOTE_SYSTEM_PROMPT)


async def _apply_note_update(user: User, note_id: int, changes: dict) -> str:
    lang = user.language
    async with SessionLocal() as session:
        note = (
            await session.execute(
                select(Note).where(Note.id == note_id, Note.user_id == user.id)
            )
        ).scalar_one_or_none()
        if note is None:
            return t("note_not_found", lang, id=note_id)

        applied: list[str] = []
        new_title = changes.get("title")
        if isinstance(new_title, str) and new_title.strip():
            note.title = new_title.strip()
            applied.append(t("change_title", lang))
        new_content = changes.get("content")
        if isinstance(new_content, str) and new_content.strip():
            if changes.get("append"):
                note.content = ((note.content or "") + "\n" + new_content.strip()).strip()
                applied.append(t("change_content_append", lang))
            else:
                note.content = new_content.strip()
                applied.append(t("change_content", lang))

        if not applied:
            return t("note_update_nothing", lang)

        note.embedding = await llm.embed(f"{note.title}\n{note.content or ''}")
        await session.commit()
        await session.refresh(note)
        title = note.title

    logger.info("Updated note %s for user %s: %s", note_id, user.id, applied)
    return t("note_updated", lang, id=note_id, title=title, changes=", ".join(applied))


async def update_note_nl(user: User, text: str) -> str:
    notes = await get_active_notes(user)
    if not notes:
        return t("note_update_none", user.language)
    parsed = await _parse_note_update(text, notes)
    try:
        note_id = int(parsed.get("note_id") or 0)
    except (TypeError, ValueError):
        note_id = 0
    if note_id <= 0:
        return t("note_update_no_target", user.language)
    return await _apply_note_update(user, note_id, parsed.get("changes") or {})


async def update_note_by_id(user: User, note_id: int, text: str) -> str:
    async with SessionLocal() as session:
        note = (
            await session.execute(
                select(Note).where(Note.id == note_id, Note.user_id == user.id)
            )
        ).scalar_one_or_none()
    if note is None:
        return t("note_not_found", user.language, id=note_id)
    parsed = await _parse_note_update(text, [note])
    return await _apply_note_update(user, note_id, parsed.get("changes") or {})


# ---- Export & attachments -------------------------------------------------

async def export_note(user: User, note_id: int, fmt: str) -> tuple[str, str] | None:
    """Create a .docx/.txt file from a note. Returns (path, filename) or None."""
    async with SessionLocal() as session:
        note = (
            await session.execute(
                select(Note).where(Note.id == note_id, Note.user_id == user.id)
            )
        ).scalar_one_or_none()
    if note is None:
        return None

    user_dir = EXPORT_ROOT / str(user.telegram_id)
    user_dir.mkdir(parents=True, exist_ok=True)
    safe = _safe_filename(note.title)
    content = note.content or ""

    if fmt == "docx":
        from docx import Document

        doc = Document()
        doc.add_heading(note.title, level=1)
        for para in content.split("\n"):
            doc.add_paragraph(para)
        path = user_dir / f"{safe}.docx"
        doc.save(str(path))
        filename = f"{safe}.docx"
    else:
        path = user_dir / f"{safe}.txt"
        path.write_text(f"{note.title}\n\n{content}", encoding="utf-8")
        filename = f"{safe}.txt"

    logger.info("Exported note %s -> %s", note_id, path)
    return str(path), filename


async def get_note_files(user: User, note_id: int) -> list[StoredFile]:
    async with SessionLocal() as session:
        stmt = (
            select(StoredFile)
            .where(StoredFile.note_id == note_id, StoredFile.user_id == user.id)
            .order_by(StoredFile.id)
        )
        return list((await session.execute(stmt)).scalars().all())


async def detach_file(user: User, file_id: int) -> str:
    """Unlink a file from its note (the file itself is kept)."""
    async with SessionLocal() as session:
        f = (
            await session.execute(
                select(StoredFile).where(
                    StoredFile.id == file_id, StoredFile.user_id == user.id
                )
            )
        ).scalar_one_or_none()
        if f is None:
            return t("file_not_found", user.language, id=file_id)
        f.note_id = None
        await session.commit()
        name = f.display_name

    logger.info("Detached file %s from its note (user %s)", file_id, user.id)
    return t("note_detached", user.language, name=name)
